import os
import django
import json

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'goathlete_backend.settings')
django.setup()

from tournaments.models import SportTemplate, SportEvent
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SPORTS_DATA = [
    {
        "name": "Cricket",
        "min_players_per_team": 2,
        "max_players_per_team": 11,
        "default_rules": {
            "format": "Limited Overs / T20",
            "default_overs": 20,
            "balls_per_over": 6,
            "scoring_units": "runs_and_wickets",
            "umpire_roles": ["Main Umpire (Bowler End)", "Leg Umpire", "Third Umpire / Scorer"]
        },
        "events": [
            {"name": "Dot Ball (0 Run)", "category": "SCORING", "points_value": 0, "requires_player": True},
            {"name": "Single (1 Run)", "category": "SCORING", "points_value": 1, "requires_player": True},
            {"name": "Two Runs (2 Runs)", "category": "SCORING", "points_value": 2, "requires_player": True},
            {"name": "Three Runs (3 Runs)", "category": "SCORING", "points_value": 3, "requires_player": True},
            {"name": "Boundary Four (4 Runs)", "category": "SCORING", "points_value": 4, "requires_player": True},
            {"name": "Sixer (6 Runs)", "category": "SCORING", "points_value": 6, "requires_player": True},
            {"name": "Wide Ball (+1 Extra)", "category": "EXTRA", "points_value": 1, "requires_player": False},
            {"name": "No Ball (+1 Extra + Free Hit)", "category": "EXTRA", "points_value": 1, "requires_player": False},
            {"name": "Leg Bye (+1 Extra)", "category": "EXTRA", "points_value": 1, "requires_player": True},
            {"name": "Bye (+1 Extra)", "category": "EXTRA", "points_value": 1, "requires_player": True},
            {"name": "Wicket (Bowled)", "category": "OUT", "points_value": 0, "requires_player": True},
            {"name": "Wicket (Caught)", "category": "OUT", "points_value": 0, "requires_player": True},
            {"name": "Wicket (LBW)", "category": "OUT", "points_value": 0, "requires_player": True},
            {"name": "Wicket (Run Out)", "category": "OUT", "points_value": 0, "requires_player": True},
            {"name": "Wicket (Stumped)", "category": "OUT", "points_value": 0, "requires_player": True},
            {"name": "Over Completed", "category": "TACTICAL", "points_value": 0, "requires_player": False},
        ]
    },
    {
        "name": "Football (Soccer)",
        "min_players_per_team": 5,
        "max_players_per_team": 11,
        "default_rules": {
            "format": "Standard / 5v5 / 7v7",
            "half_duration_minutes": 45,
            "halves": 2,
            "scoring_units": "goals",
            "umpire_roles": ["Head Referee", "Assistant Referee 1", "Assistant Referee 2"]
        },
        "events": [
            {"name": "Goal", "category": "SCORING", "points_value": 1, "requires_player": True},
            {"name": "Penalty Goal", "category": "SCORING", "points_value": 1, "requires_player": True},
            {"name": "Own Goal (Opponent Point)", "category": "SCORING", "points_value": 1, "requires_player": True},
            {"name": "Yellow Card", "category": "PENALTY", "points_value": 0, "requires_player": True},
            {"name": "Red Card (Send Off)", "category": "PENALTY", "points_value": 0, "requires_player": True},
            {"name": "Corner Kick", "category": "TACTICAL", "points_value": 0, "requires_player": True},
            {"name": "Offside", "category": "TACTICAL", "points_value": 0, "requires_player": True},
            {"name": "Foul Committed", "category": "PENALTY", "points_value": 0, "requires_player": True},
            {"name": "Substitution", "category": "TACTICAL", "points_value": 0, "requires_player": True},
        ]
    },
    {
        "name": "Badminton",
        "min_players_per_team": 1,
        "max_players_per_team": 2,
        "default_rules": {
            "format": "Singles / Doubles",
            "points_to_win_set": 21,
            "max_points_deuce": 30,
            "best_of_sets": 3,
            "scoring_units": "points_and_sets",
            "umpire_roles": ["Umpire", "Service Judge", "Line Judge"]
        },
        "events": [
            {"name": "Rally Point", "category": "SCORING", "points_value": 1, "requires_player": True},
            {"name": "Smash Winner", "category": "SCORING", "points_value": 1, "requires_player": True},
            {"name": "Service Ace", "category": "SCORING", "points_value": 1, "requires_player": True},
            {"name": "Fault (Net / Out)", "category": "PENALTY", "points_value": 0, "requires_player": True},
            {"name": "Service Fault", "category": "PENALTY", "points_value": 0, "requires_player": True},
            {"name": "Set Won", "category": "MILESTONE", "points_value": 0, "requires_player": False},
        ]
    },
    {
        "name": "Basketball",
        "min_players_per_team": 3,
        "max_players_per_team": 5,
        "default_rules": {
            "format": "3x3 / Full Court 5v5",
            "quarters": 4,
            "quarter_duration_minutes": 10,
            "scoring_units": "points",
            "umpire_roles": ["Lead Referee", "Trail Referee", "Table Official / Scorer"]
        },
        "events": [
            {"name": "Free Throw (1 Point)", "category": "SCORING", "points_value": 1, "requires_player": True},
            {"name": "Field Goal (2 Points)", "category": "SCORING", "points_value": 2, "requires_player": True},
            {"name": "Three-Pointer (3 Points)", "category": "SCORING", "points_value": 3, "requires_player": True},
            {"name": "Personal Foul", "category": "PENALTY", "points_value": 0, "requires_player": True},
            {"name": "Technical Foul", "category": "PENALTY", "points_value": 0, "requires_player": True},
            {"name": "Timeout Called", "category": "TACTICAL", "points_value": 0, "requires_player": False},
            {"name": "Turnover", "category": "TACTICAL", "points_value": 0, "requires_player": True},
        ]
    },
    {
        "name": "Tennis",
        "min_players_per_team": 1,
        "max_players_per_team": 2,
        "default_rules": {
            "format": "Singles / Doubles",
            "games_to_win_set": 6,
            "tiebreak_points": 7,
            "best_of_sets": 3,
            "scoring_units": "points_games_sets",
            "umpire_roles": ["Chair Umpire", "Line Umpire"]
        },
        "events": [
            {"name": "Point Won (15/30/40/Game)", "category": "SCORING", "points_value": 1, "requires_player": True},
            {"name": "Ace", "category": "SCORING", "points_value": 1, "requires_player": True},
            {"name": "Double Fault", "category": "PENALTY", "points_value": 0, "requires_player": True},
            {"name": "Break Point Converted", "category": "SCORING", "points_value": 1, "requires_player": True},
            {"name": "Game Won", "category": "MILESTONE", "points_value": 0, "requires_player": False},
            {"name": "Set Won", "category": "MILESTONE", "points_value": 0, "requires_player": False},
        ]
    },
    {
        "name": "Volleyball",
        "min_players_per_team": 2,
        "max_players_per_team": 6,
        "default_rules": {
            "format": "Indoor 6v6 / Beach 2v2",
            "points_to_win_set": 25,
            "deciding_set_points": 15,
            "best_of_sets": 5,
            "scoring_units": "points_and_sets",
            "umpire_roles": ["First Referee", "Second Referee", "Scorer"]
        },
        "events": [
            {"name": "Spike Kill Point", "category": "SCORING", "points_value": 1, "requires_player": True},
            {"name": "Service Ace Point", "category": "SCORING", "points_value": 1, "requires_player": True},
            {"name": "Block Point", "category": "SCORING", "points_value": 1, "requires_player": True},
            {"name": "Opponent Attack Error", "category": "SCORING", "points_value": 1, "requires_player": False},
            {"name": "Service Error", "category": "PENALTY", "points_value": 0, "requires_player": True},
            {"name": "Net Touch Violation", "category": "PENALTY", "points_value": 0, "requires_player": True},
            {"name": "Set Won", "category": "MILESTONE", "points_value": 0, "requires_player": False},
        ]
    },
    {
        "name": "Table Tennis",
        "min_players_per_team": 1,
        "max_players_per_team": 2,
        "default_rules": {
            "format": "Singles / Doubles",
            "points_to_win_game": 11,
            "must_win_by": 2,
            "best_of_games": 5,
            "scoring_units": "points_and_games",
            "umpire_roles": ["Match Umpire", "Assistant Umpire"]
        },
        "events": [
            {"name": "Point Won", "category": "SCORING", "points_value": 1, "requires_player": True},
            {"name": "Service Winner", "category": "SCORING", "points_value": 1, "requires_player": True},
            {"name": "Edge Ball", "category": "SCORING", "points_value": 1, "requires_player": True},
            {"name": "Service Fault", "category": "PENALTY", "points_value": 0, "requires_player": True},
            {"name": "Game Won", "category": "MILESTONE", "points_value": 0, "requires_player": False},
        ]
    },
    {
        "name": "Kabaddi",
        "min_players_per_team": 7,
        "max_players_per_team": 7,
        "default_rules": {
            "format": "Standard 7v7",
            "half_duration_minutes": 20,
            "halves": 2,
            "scoring_units": "points",
            "umpire_roles": ["Referee", "Umpire 1", "Umpire 2", "Scorer"]
        },
        "events": [
            {"name": "Touch Point (Raid)", "category": "SCORING", "points_value": 1, "requires_player": True},
            {"name": "Bonus Point (Raid)", "category": "SCORING", "points_value": 1, "requires_player": True},
            {"name": "Tackle Point (Defense)", "category": "SCORING", "points_value": 1, "requires_player": True},
            {"name": "Super Tackle (2 Points)", "category": "SCORING", "points_value": 2, "requires_player": True},
            {"name": "Super Raid (3+ Points)", "category": "SCORING", "points_value": 3, "requires_player": True},
            {"name": "All Out (+2 Points)", "category": "SCORING", "points_value": 2, "requires_player": False},
            {"name": "Do-or-Die Raid Out", "category": "PENALTY", "points_value": 0, "requires_player": True},
        ]
    },
    {
        "name": "Pickleball",
        "min_players_per_team": 1,
        "max_players_per_team": 2,
        "default_rules": {
            "format": "Singles / Doubles",
            "points_to_win_game": 11,
            "must_win_by": 2,
            "best_of_games": 3,
            "scoring_units": "points_and_games",
            "umpire_roles": ["Head Referee", "Line Judge"]
        },
        "events": [
            {"name": "Point Won (Serving Team)", "category": "SCORING", "points_value": 1, "requires_player": True},
            {"name": "Side Out (Service Lost)", "category": "TACTICAL", "points_value": 0, "requires_player": False},
            {"name": "Kitchen / Non-Volley Zone Fault", "category": "PENALTY", "points_value": 0, "requires_player": True},
            {"name": "Game Won", "category": "MILESTONE", "points_value": 0, "requires_player": False},
        ]
    },
    {
        "name": "Box / Turf Cricket",
        "min_players_per_team": 4,
        "max_players_per_team": 8,
        "default_rules": {
            "format": "Short Format Box Cricket",
            "default_overs": 6,
            "balls_per_over": 6,
            "scoring_units": "runs_and_wickets",
            "box_rules": "Direct roof hit out, side wall 1/2 runs",
            "umpire_roles": ["Turf Referee / Scorer"]
        },
        "events": [
            {"name": "Single (1 Run)", "category": "SCORING", "points_value": 1, "requires_player": True},
            {"name": "Two Runs (2 Runs)", "category": "SCORING", "points_value": 2, "requires_player": True},
            {"name": "Wall Boundary (4 Runs)", "category": "SCORING", "points_value": 4, "requires_player": True},
            {"name": "Direct Sixer (6 Runs)", "category": "SCORING", "points_value": 6, "requires_player": True},
            {"name": "Direct Net / Wall Catch Out", "category": "OUT", "points_value": 0, "requires_player": True},
            {"name": "Wicket (Bowled / Run Out)", "category": "OUT", "points_value": 0, "requires_player": True},
            {"name": "Wide / No Ball (+1 Extra)", "category": "EXTRA", "points_value": 1, "requires_player": False},
        ]
    }
]

def seed_database():
    print("--- Seeding Sports Templates and Events ---")
    for sdata in SPORTS_DATA:
        template, created = SportTemplate.objects.update_or_create(
            name=sdata["name"],
            defaults={
                "min_players_per_team": sdata["min_players_per_team"],
                "max_players_per_team": sdata["max_players_per_team"],
                "is_active": True,
                "default_rules": sdata["default_rules"],
            }
        )
        print(f"[{'CREATED' if created else 'UPDATED'}] Sport Template: {template.name}")

        # Clear existing and re-create events
        template.events.all().delete()
        for edata in sdata["events"]:
            event = SportEvent.objects.create(
                sport_template=template,
                name=edata["name"],
                category=edata["category"],
                points_value=edata["points_value"],
                requires_player=edata["requires_player"]
            )
        print(f"  -> Added {len(sdata['events'])} scoring events.")

def generate_excel_matrix():
    wb = openpyxl.Workbook()
    
    # Header styling
    header_fill = PatternFill(start_color="E65100", end_color="E65100", fill_type="solid") # GoAthlete Orange
    sub_header_fill = PatternFill(start_color="FF9800", end_color="FF9800", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Arial", size=10, bold=True)
    normal_font = Font(name="Arial", size=10)
    thin_border = Border(
        left=Side(style='thin', color='D0D0D0'),
        right=Side(style='thin', color='D0D0D0'),
        top=Side(style='thin', color='D0D0D0'),
        bottom=Side(style='thin', color='D0D0D0')
    )
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    # Sheet 1: Sports Master & Rules
    ws_sports = wb.active
    ws_sports.title = "Sports Configurations"

    sports_headers = [
        "Sport ID", "Sport Name", "Min Players", "Max Players", 
        "Default Format", "Scoring Unit", "Match Duration / Overs", "Umpire Roles", "JSON Config"
    ]
    
    ws_sports.append(sports_headers)
    for col_num in range(1, len(sports_headers) + 1):
        cell = ws_sports.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    for idx, s in enumerate(SPORTS_DATA, start=1):
        rules = s["default_rules"]
        format_val = rules.get("format", "")
        unit_val = rules.get("scoring_units", "")
        duration_val = str(rules.get("default_overs", rules.get("half_duration_minutes", rules.get("points_to_win_set", rules.get("quarter_duration_minutes", "")))))
        roles_val = ", ".join(rules.get("umpire_roles", ["Match Umpire"]))
        json_val = json.dumps(rules)

        ws_sports.append([
            idx, s["name"], s["min_players_per_team"], s["max_players_per_team"],
            format_val, unit_val, duration_val, roles_val, json_val
        ])

    for row in ws_sports.iter_rows(min_row=2, max_row=len(SPORTS_DATA)+1):
        for cell in row:
            cell.font = normal_font
            cell.border = thin_border
            if cell.column in [1, 3, 4, 7]:
                cell.alignment = center_align

    # Sheet 2: Sport Events & Scoring Actions
    ws_events = wb.create_sheet(title="Scoring Events & Actions")
    events_headers = [
        "Event ID", "Sport Name", "Action / Event Name", "Category", 
        "Points Value", "Requires Player Selection", "Umpire Action Trigger"
    ]
    ws_events.append(events_headers)
    for col_num in range(1, len(events_headers) + 1):
        cell = ws_events.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    event_id = 1
    for s in SPORTS_DATA:
        for e in s["events"]:
            ws_events.append([
                event_id,
                s["name"],
                e["name"],
                e["category"],
                e["points_value"],
                "Yes" if e["requires_player"] else "No",
                f"Umpire 1-Tap button: [{e['name']}]"
            ])
            event_id += 1

    for row in ws_events.iter_rows(min_row=2, max_row=event_id):
        for cell in row:
            cell.font = normal_font
            cell.border = thin_border
            if cell.column in [1, 4, 5, 6]:
                cell.alignment = center_align

    # Sheet 3: Role Permission Matrix
    ws_roles = wb.create_sheet(title="Role Permissions Matrix")
    role_headers = ["Screen / Feature", "Umpire / Scorer", "Tournament Organizer", "Team Captain", "Player / Spectator"]
    ws_roles.append(role_headers)
    for col_num in range(1, len(role_headers) + 1):
        cell = ws_roles.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    role_rows = [
        ["Create Tournament & Schedule Fixtures", "View Only", "Full Access (Create/Edit)", "View Only", "View Only"],
        ["Register Teams & Add Players", "View Only", "Full Access", "Full Access (Own Team)", "View Only"],
        ["Scan Athlete QR Code / Lookup GOATH-ID", "Full Access", "Full Access", "Full Access", "View Only (Own Profile)"],
        ["Start Match (Scheduled -> Ongoing)", "Full Access", "Full Access", "No Access", "No Access"],
        ["Live Scoring Panel (Log Runs, Goals, Points)", "Full Access (Edit)", "Full Access (Edit)", "Read-Only (View Score)", "Read-Only (View Score)"],
        ["Undo Score / Event", "Full Access", "Full Access", "No Access", "No Access"],
        ["End Match & Declare Winner", "Full Access", "Full Access", "No Access", "No Access"],
        ["View Live Broadcast Scoreboard", "Full Access", "Full Access", "Full Access", "Full Access (Auto-Refreshing)"],
    ]

    for r in role_rows:
        ws_roles.append(r)

    for row in ws_roles.iter_rows(min_row=2, max_row=len(role_rows)+1):
        for cell in row:
            cell.font = normal_font
            cell.border = thin_border
            if "Full Access" in str(cell.value):
                cell.font = Font(name="Arial", size=10, bold=True, color="2E7D32") # Dark green
            elif "No Access" in str(cell.value):
                cell.font = Font(name="Arial", size=10, color="C62828") # Dark red

    # Auto-adjust column widths
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

    excel_path = "sports_configurations_matrix.xlsx"
    wb.save(excel_path)
    print(f"--- Excel Matrix saved to {excel_path} ---")

if __name__ == "__main__":
    seed_database()
    generate_excel_matrix()
